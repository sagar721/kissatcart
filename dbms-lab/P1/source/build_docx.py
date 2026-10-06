"""Builds the editable Word version of the P1 record from the same HTML source
used for the PDF (source/P1_DBMS_Lab_Record.html, produced by build.py).

Run after build.py:  python3 source/build_docx.py
Output: P1_DBMS_Lab_Record_Sagar_Kumar.docx
"""
import re
import subprocess
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC_HTML = ROOT / "source" / "P1_DBMS_Lab_Record.html"
OUT = ROOT / "P1_DBMS_Lab_Record_Sagar_Kumar.docx"
CHROMIUM = "/opt/pw-browsers/chromium"
SERIF, MONO = "Times New Roman", "Consolas"
CONTENT_MM = 174


# ----------------------------------------------------------------- helpers
def svg_to_png(svg: Path) -> Path:
    """Rasterise a diagram SVG at 2x for embedding in Word."""
    m = re.search(r'viewBox="0 0 (\d+) (\d+)"', svg.read_text())
    w, h = int(m.group(1)), int(m.group(2))
    png = svg.with_suffix(".png")
    subprocess.run([CHROMIUM, "--headless", "--no-sandbox", "--disable-gpu", "--hide-scrollbars",
                    "--force-device-scale-factor=2", f"--window-size={w},{h + 100}",
                    f"--screenshot={png}", svg.as_uri()],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    subprocess.run(["convert", str(png), "-crop", f"{2 * w}x{2 * h}+0+0", "+repage", str(png)], check=True)
    return png


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_font(run, name=SERIF, size=None, bold=None, italic=None, underline=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline


def add_field(paragraph, instr):
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
            run._r.append(el)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
            run._r.append(el)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
    set_font(run, size=9)


def para_spacing(p, before=0, after=6, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def inline(p, node, size=11.5, bold=False, italic=False, underline=False, mono=False):
    """Add the inline content of an HTML node to a docx paragraph."""
    for ch in node.children:
        if isinstance(ch, NavigableString):
            text = re.sub(r"\s+", " ", str(ch))
            if not text:
                continue
            r = p.add_run(text)
            set_font(r, MONO if mono else SERIF, (size - 2) if mono else size, bold, italic, underline)
        elif isinstance(ch, Tag):
            cls = ch.get("class") or []
            if ch.name == "br":
                p.add_run().add_break()
            else:
                inline(p, ch, size,
                       bold or ch.name in ("b", "strong") or "u" in cls,
                       italic or ch.name in ("i", "em") or "fk" in cls,
                       underline or "u" in cls,
                       mono or ch.name == "code")


def trim_paragraph(p):
    if p.runs:
        p.runs[0].text = p.runs[0].text.lstrip()
        p.runs[-1].text = p.runs[-1].text.rstrip()


# ----------------------------------------------------------------- document
doc = Document()
st = doc.styles["Normal"]
st.font.name = SERIF
st.font.size = Pt(11.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)

sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = Mm(18)
sec.top_margin, sec.bottom_margin = Mm(22), Mm(20)
sec.header_distance = sec.footer_distance = Mm(9)
sec.different_first_page_header_footer = True

hp = sec.header.paragraphs[0]
hp.text = ""
hp.style = doc.styles["Normal"]
r = hp.add_run("U25AMILPC307 – DBMS Lab  |  Practical P1\tSagar Kumar  |  PRN: SOE25BTAM29  |  Section A")
set_font(r, size=9)
hp.paragraph_format.tab_stops.add_tab_stop(Mm(CONTENT_MM), alignment=2)
fp = sec.footer.paragraphs[0]
fp.style = doc.styles["Normal"]
fp.paragraph_format.tab_stops.add_tab_stop(Mm(CONTENT_MM), alignment=2)
r = fp.add_run("College Management System – ER Modeling and Database Creation\tPage ")
set_font(r, size=9)
add_field(fp, "PAGE")
r = fp.add_run(" of ")
set_font(r, size=9)
add_field(fp, "NUMPAGES")

soup = BeautifulSoup(SRC_HTML.read_text(), "html.parser")

# ---- cover page (framed with a one-cell table)
cover = soup.select_one("section.cover")
frame = doc.add_table(rows=1, cols=1)
frame.style = "Table Grid"
frame.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = frame.rows[0].cells[0]
cell.width = Mm(CONTENT_MM)
first = cell.paragraphs[0]


def cover_line(text, size, bold=False, before=0, after=6, spacing=None, first_par=None):
    p = first_par or cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para_spacing(p, before, after)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    if spacing:
        rPr = r._element.get_or_add_rPr()
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(spacing))
        rPr.append(sp)
    return p


cover_line("PRACTICAL RECORD", 13, before=40, after=14, spacing=60, first_par=first)
cover_line("DBMS Lab", 24, bold=True, after=4)
cover_line("Course Code: U25AMILPC307", 13, after=24)
cover_line("PRACTICAL P1", 15, bold=True, after=10, spacing=20)
cover_line(cover.select_one(".ptitle").get_text(" ", strip=True), 14, after=10)
p = cell.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
para_spacing(p, 0, 36)
set_font(p.add_run("Case Study: "), size=12)
set_font(p.add_run("College Management System"), size=12, bold=True)
cover_line("STUDENT DETAILS", 12.5, bold=True, after=8, spacing=20)
details = cell.add_table(rows=0, cols=2)
details.style = "Table Grid"
details.alignment = WD_TABLE_ALIGNMENT.CENTER
for tr in cover.select("table.details tr"):
    row = details.add_row().cells
    for c, src, b in ((row[0], tr.th, True), (row[1], tr.td, False)):
        c.width = Mm(50 if b else 80)
        cp = c.paragraphs[0]
        para_spacing(cp, 3, 3)
        set_font(cp.add_run(src.get_text(strip=True)), size=12, bold=b)
        if b:
            shade(c, "EFEFEF")
cover_line("", 12, after=60)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---- body
def add_table(tbl_tag, font=10.5):
    rows = tbl_tag.find_all("tr")
    ncols = max(sum(int(c.get("colspan", 1)) for c in r.find_all(["td", "th"])) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = "Table Grid"
    occupied = set()
    for ri, tr in enumerate(rows):
        ci = 0
        for c in tr.find_all(["td", "th"]):
            while (ri, ci) in occupied:
                ci += 1
            rs, cs = int(c.get("rowspan", 1)), int(c.get("colspan", 1))
            target = t.cell(ri, ci)
            if rs > 1 or cs > 1:
                target = target.merge(t.cell(ri + rs - 1, ci + cs - 1))
            for dr in range(rs):
                for dc in range(cs):
                    occupied.add((ri + dr, ci + dc))
            p = target.paragraphs[0]
            para_spacing(p, 1, 1, 1.0)
            is_head = c.name == "th"
            inline(p, c, size=font, bold=is_head, underline="u" in (c.get("class") or []),
                   italic="fk" in (c.get("class") or []))
            trim_paragraph(p)
            if is_head:
                shade(target, "E9E9E9")
            ci += cs
    # column widths from the header row's "width:NN%" styles; the rest share what is left
    head = rows[0].find_all(["td", "th"])
    pct = [None] * ncols
    for i, c in enumerate(head[:ncols]):
        m = re.search(r"width:\s*([\d.]+)%", c.get("style", ""))
        if m and int(c.get("colspan", 1)) == 1:
            pct[i] = float(m.group(1))
    if any(pct):
        free = [i for i, v in enumerate(pct) if v is None]
        left = max(100 - sum(v for v in pct if v), 5)
        for i in free:
            pct[i] = left / len(free)
        t.autofit = False
        for gc, v in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), pct):
            gc.set(qn("w:w"), str(int(CONTENT_MM * v / 100 * 56.7)))
        for row in t.rows:
            for i, c in enumerate(row.cells[:ncols]):
                c.width = Mm(CONTENT_MM * pct[i] / 100)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def boxed(text_lines, mono=True, size=8.5, fill="F6F6F6"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    shade(c, fill)
    p = c.paragraphs[0]
    para_spacing(p, 2, 2, 1.0)
    for i, line in enumerate(text_lines):
        if i:
            p.add_run().add_break()
        set_font(p.add_run(line), MONO if mono else SERIF, size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def heading(text, level, page_break=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.keep_with_next = True
    pf.page_break_before = page_break
    para_spacing(p, 14 if level == 1 else 8, 5)
    set_font(p.add_run(text), size=14 if level == 1 else 12, bold=True)
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", "000000")):
            b.set(qn(k), v)
        bdr.append(b)
        pPr.append(bdr)


def figure(fig):
    img = fig.find("img")
    src = (ROOT / "source" / img["src"]).resolve()
    if src.suffix == ".svg":
        src = svg_to_png(src)
        width = Mm(CONTENT_MM * (0.92 if "92%" in img.get("style", "") else 1.0))
    else:
        m = re.search(r"width:([\d.]+)mm", img.get("style", ""))
        width = Mm(float(m.group(1)))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    para_spacing(p, 6, 2)
    p.add_run().add_picture(str(src), width=width)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para_spacing(cap, 0, 10)
    inline(cap, fig.find("figcaption"), size=10, italic=True)
    trim_paragraph(cap)


for node in soup.body.children:
    if not isinstance(node, Tag) or node.name == "section":
        continue
    cls = node.get("class") or []
    if node.name == "h2":
        heading(node.get_text(" ", strip=True), 1, "page-break" in cls)
    elif node.name == "h3":
        heading(node.get_text(" ", strip=True), 2)
    elif node.name == "p":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para_spacing(p, 0, 6)
        inline(p, node)
        trim_paragraph(p)
    elif node.name in ("ul", "ol"):
        style = "List Bullet" if node.name == "ul" else "List Number"
        for li in node.find_all("li", recursive=False):
            p = doc.add_paragraph(style=style)
            para_spacing(p, 0, 3)
            inline(p, li)
            trim_paragraph(p)
    elif node.name == "table":
        add_table(node, 10 if "avoid" in cls else 10.5)
    elif node.name == "pre":
        boxed(node.get_text().rstrip("\n").split("\n"))
    elif node.name == "figure":
        figure(node)
    elif node.name == "div" and "note" in cls:
        boxed([node.get_text(" ", strip=True)], mono=False, size=10.3, fill="F2F2F2")
    elif node.name == "div" and "schema" in cls:
        for sp in node.find_all("p"):
            p = doc.add_paragraph()
            para_spacing(p, 0, 4)
            inline(p, sp, size=11.5, mono=True)
            trim_paragraph(p)
    elif node.name == "div" and "sign" in cls:
        p = doc.add_paragraph()
        para_spacing(p, 26, 6)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(CONTENT_MM), alignment=2)
        divs = node.find_all("div")
        set_font(p.add_run(divs[0].get_text(strip=True) + "\t" + divs[1].get_text(strip=True)), size=11)
    elif node.name == "dl":
        for item in node.find_all(["dt", "dd"]):
            p = doc.add_paragraph()
            if item.name == "dt":
                para_spacing(p, 7, 2)
                p.paragraph_format.keep_with_next = True
                inline(p, item, bold=True)
            else:
                para_spacing(p, 0, 4)
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                inline(p, item)
            trim_paragraph(p)

doc.core_properties.title = "DBMS Lab Practical P1 – Database Creation and ER Modeling (College Management System)"
doc.core_properties.author = "Sagar Kumar (SOE25BTAM29)"
doc.save(OUT)
print("wrote", OUT)
